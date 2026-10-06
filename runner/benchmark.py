import json, re, resource, subprocess, sys, tempfile, time
from pathlib import Path
import docx, pypdf, pypdfium2, pytesseract

FX = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures")
HEADS = ("Summary", "Skills", "Experience", "Education", "Projects")
MISMATCH = 0.10

def rss(who): return resource.getrusage(who).ru_maxrss / 1024

def native(p):
    if p.suffix == ".docx":
        return "\n".join(x.text for x in docx.Document(p).paragraphs)
    return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(p).pages)

def ocr(p):
    if p.suffix == ".docx":
        d = tempfile.mkdtemp()
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, str(p)], check=True, capture_output=True)
        p = Path(d) / (p.stem + ".pdf")
    pdf = pypdfium2.PdfDocument(str(p))
    return "\n".join(pytesseract.image_to_string(pdf[i].render(scale=2).to_pil()) for i in range(len(pdf)))

def words(t): return set(re.findall(r"[a-z0-9]{3,}", t.lower()))

def parse(t):
    L = [l.strip() for l in t.splitlines() if l.strip()]
    sec, cur = {}, None
    for l in L:
        if l in HEADS: cur = l; sec[cur] = []
        elif cur: sec[cur].append(l)
    m = re.search(r"[\w.]+@[\w.]+", t)
    return dict(name=L[0] if L else "", email=m.group(0) if m else "",
        skills=[s.strip() for s in " ".join(sec.get("Skills", [])).split(",") if s.strip()],
        experience=[l for l in sec.get("Experience", []) if re.search(r"\(.+\)$", l) and not l.startswith("•")])

def expected(lab):
    return dict(name=lab["name"], email=lab["email"], skills=lab["skills"],
        experience=[f'{e["title"]}, {e["org"]} ({e["dates"]})' for e in lab["experience"]])

def score(got, exp, tot):
    for k in ("name", "email"):
        tot[k][0] += got[k].lower() == exp[k].lower(); tot[k][1] += 1; tot[k][2] += 1
    for k in ("skills", "experience"):
        g, e = {x.lower() for x in got[k]}, {x.lower() for x in exp[k]}
        tot[k][0] += len(g & e); tot[k][1] += len(g); tot[k][2] += len(e)

labels = json.load(open(FX / "labels.json"))
files = [str(FX / l["file"]) for l in labels]

t0 = time.time()
try:
    r = subprocess.run(["clamscan", "--no-summary"] + files, capture_output=True, text=True)
    scan = f"{time.time() - t0:.1f}s, {rss(resource.RUSAGE_CHILDREN):.0f} MB peak, infected lines: {r.stdout.count('FOUND')}"
except FileNotFoundError:
    scan = "clamscan not installed"

tot = {k: [0, 0, 0] for k in ("name", "email", "skills", "experience")}
rows = []
for lab in labels:
    p = FX / lab["file"]
    row = dict(id=lab["id"], expected=lab["expected_status"])
    try:
        t = time.time(); nat = native(p); row["native_s"] = round(time.time() - t, 2)
        t = time.time(); ocrt = ocr(p); row["ocr_s"] = round(time.time() - t, 2)
        nw, ow = words(nat), words(ocrt)
        row["missing_in_ocr"] = round(len(nw - ow) / len(nw), 2) if nw else 1.0
        text = nat if nw else ocrt
        row["status"] = "needs_clarification" if nw and row["missing_in_ocr"] > MISMATCH else "ready"
        if row["status"] == "ready" and lab["expected_status"] == "ready": score(parse(text), expected(lab), tot)
    except Exception as e:
        row["status"] = "failed_unreadable"; row["error"] = type(e).__name__
    row["ok"] = row["status"] == row["expected"]
    rows.append(row)

for r in rows: print(r)
print("clamav:", scan)
for k, (tp, g, e) in tot.items():
    print(f"{k}: precision {tp / g if g else 0:.2f} recall {tp / e if e else 0:.2f}")
print(f"status correct: {sum(r['ok'] for r in rows)}/{len(rows)}")
print(f"peak memory: python {rss(resource.RUSAGE_SELF):.0f} MB, child tools {rss(resource.RUSAGE_CHILDREN):.0f} MB")
Path("runner/results").mkdir(parents=True, exist_ok=True)
json.dump(dict(rows=rows, totals=tot, clamav=scan), open("runner/results/rules_only.json", "w"), indent=1)
